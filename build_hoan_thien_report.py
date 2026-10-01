# -*- coding: utf-8 -*-
"""
KỊCH BẢN TỰ ĐỘNG XÂY DỰNG BÁO CÁO ĐỒ ÁN KỲ HOÀN THIỆN
Đề tài: XÂY DỰNG WEBSITE THƯƠNG MẠI ĐIỆN TỬ BÁN ĐỒ ĂN TRỰC TUYẾN FOODVD
        HỖ TRỢ TRẢI NGHIỆM 3D VÀ QUẢN LÝ VẬN HÀNH ĐA PHÂN HỆ
Sinh viên thực hiện: Vũ Dũng - MSSV: 52310079 - Lớp: 523100B
Trường Đại học Phương Đông - Khoa Công nghệ Số và Truyền thông
"""

from __future__ import annotations
import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

BASE = Path(r"c:\Users\Admin\Desktop\523100B\DOANTOTNGHIEP")
DIAGRAMS_DIR = BASE / "diagrams"
OUT = BASE / "BAO_CAO_DO_AN_KY_FOODVD_HOAN_THIEN.docx"

# ================= FORMATTING HELPERS =================
def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_border(cell, color: str = "D9D9D9") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)

def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    values = {"top": top, "start": start, "bottom": bottom, "end": end}
    for key, value in values.items():
        tag = f"w:{key}"
        node = margins.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_font(run, name: str = "Times New Roman", size: int | None = None, bold: bool | None = None, italic: bool | None = None, color: RGBColor | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if color:
        run.font.color.rgb = color
    else:
        run.font.color.rgb = RGBColor(0, 0, 0)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic

def add_paragraph(doc: Document, text: str = "", style: str | None = None, bold_lead: str | None = None, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_font(lead, bold=True, size=13)
        rest = p.add_run(text[len(bold_lead):])
        set_font(rest, size=13)
    else:
        r = p.add_run(text)
        set_font(r, size=13)
    return p

def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    color = RGBColor(15, 23, 42) if level == 1 else RGBColor(30, 41, 59)
    size = 16 if level == 1 else (14 if level == 2 else 13)
    for run in p.runs:
        set_font(run, bold=True, size=size, color=color)
    return p

def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.3
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(item)
        set_font(r, size=13)

def add_figure(doc: Document, image_path: Path | str, caption: str, width=Inches(6.0)):
    img_p = Path(image_path)
    if img_p.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(str(img_p), width=width)

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(10)
        run_cap = p_cap.add_run(caption)
        set_font(run_cap, size=11, bold=True, italic=True, color=RGBColor(51, 65, 85))
    else:
        print(f"Warning: Figure not found at {img_p}")

def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_shading(hdr_cells[i], "1E293B")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, start=120, end=120)
        set_cell_border(hdr_cells[i], "0F172A")
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            set_font(r, size=11, bold=True, color=RGBColor(255, 255, 255))

    # Data Rows
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_shading(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, start=100, end=100)
            set_cell_border(row_cells[c_idx], "CBD5E1")
            row_cells[c_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                set_font(r, size=10.5)

    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

print("Formatting helpers loaded successfully.")
