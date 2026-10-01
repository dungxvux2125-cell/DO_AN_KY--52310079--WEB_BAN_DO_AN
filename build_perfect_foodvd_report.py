# -*- coding: utf-8 -*-
"""
HỆ THỐNG XÂY DỰNG BÁO CÁO ĐỒ ÁN KỲ FOODVD CHUẨN MỰC TUYỆT ĐỐI
Đề tài: XÂY DỰNG WEBSITE THƯƠNG MẠI ĐIỆN TỬ BÁN ĐỒ ĂN TRỰC TUYẾN FOODVD
        HỖ TRỢ TRẢI NGHIỆM 3D VÀ QUẢN LÝ VẬN HÀNH ĐA PHÂN HỆ
Sinh viên: Vũ Dũng - MSSV: 52310079 - Lớp: 523100B
Trường Đại học Phương Đông - Khoa Công nghệ Số và Truyền thông
"""

import sys
from pathlib import Path

# Configure utf-8 stdout for Windows
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

BASE = Path(r"c:\Users\Admin\Desktop\523100B\DOANTOTNGHIEP")
DIAGRAMS = BASE / "diagrams"
OUT_PATH = BASE / "BAO_CAO_DO_AN_KY_FOODVD_HOAN_THIEN.docx"

doc = Document()

# Page Setup: A4, Left 3.0cm, Right 2.0cm, Top 2.0cm, Bottom 2.0cm
for sec in doc.sections:
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)

# ================= HELPER FUNCTIONS =================
def set_cell_border(cell, color="000000", sz="4", val="single"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), val)
        element.set(qn("w:sz"), sz)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)

def set_cell_margins(cell, top=70, start=100, bottom=70, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    values = {"top": top, "start": start, "bottom": bottom, "end": end}
    for key, value in values.items():
        tag = f"w:{key}"
        node = margins.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_font(run, name="Times New Roman", size=13, bold=None, italic=None, color=RGBColor(0, 0, 0)):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.color.rgb = color
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic

def p(text="", bold_lead=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4, space_before=0, italic=False, bold=False):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(space_after)
    par.paragraph_format.space_before = Pt(space_before)
    par.paragraph_format.line_spacing = 1.35
    par.alignment = align
    if bold_lead and text.startswith(bold_lead):
        r1 = par.add_run(bold_lead)
        set_font(r1, bold=True, size=13)
        r2 = par.add_run(text[len(bold_lead):])
        set_font(r2, size=13, italic=italic)
    else:
        r = par.add_run(text)
        set_font(r, size=13, italic=italic, bold=bold)
    return par

def h(text, level=1):
    par = doc.add_heading(text, level=level)
    par.paragraph_format.keep_with_next = True
    par.paragraph_format.space_before = Pt(12 if level == 1 else (8 if level == 2 else 6))
    par.paragraph_format.space_after = Pt(4)
    par.paragraph_format.line_spacing = 1.25
    size = 15 if level == 1 else (13.5 if level == 2 else 13)
    for r in par.runs:
        set_font(r, bold=True, size=size, color=RGBColor(0, 0, 0))
    return par

def bullets(items: list[str]):
    for it in items:
        par = doc.add_paragraph(style="List Bullet")
        par.paragraph_format.space_after = Pt(2)
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.line_spacing = 1.3
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = par.add_run(it)
        set_font(r, size=13)

def fig(filename, caption, width=Inches(5.8)):
    path = DIAGRAMS / filename
    if path.exists():
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(6)
        par.paragraph_format.space_after = Pt(2)
        run = par.add_run()
        run.add_picture(str(path), width=width)

        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(1)
        cap.paragraph_format.space_after = Pt(8)
        rc = cap.add_run(caption)
        set_font(rc, size=11, bold=True, italic=True, color=RGBColor(0, 0, 0))
    else:
        print(f"Warning: Figure not found: {filename}")

# Chuẩn bảng biểu: Chữ đen, Nền trắng, Viền đen chuẩn học thuật
def tbl(headers, rows, widths=None):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row (Nền trắng, Chữ đen đậm)
    for i, title in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = title
        set_cell_border(cell, color="000000", sz="6")
        set_cell_margins(cell, top=80, bottom=80, start=90, end=90)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        par = cell.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(2)
        par.paragraph_format.space_before = Pt(0)
        for r in par.runs:
            set_font(r, size=11, bold=True, color=RGBColor(0, 0, 0))

    # Data Rows (Nền trắng, Chữ đen thường)
    for r_idx, r_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        for c_idx, val in enumerate(r_data):
            cell = row.cells[c_idx]
            cell.text = str(val)
            set_cell_border(cell, color="666666", sz="4")
            set_cell_margins(cell, top=60, bottom=60, start=80, end=80)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            par = cell.paragraphs[0]
            par.paragraph_format.line_spacing = 1.2
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.space_before = Pt(0)
            if c_idx == 0:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in par.runs:
                set_font(r, size=10.5, color=RGBColor(0, 0, 0))

    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)

    par_after = doc.add_paragraph()
    par_after.paragraph_format.space_after = Pt(4)

# Bảng mục lục không viền (đảm bảo căn số trang thẳng tắp)
def toc_table(entries: list[tuple[str, str, int]]):
    # entries: [(title, page_num, indent_level)]
    table = doc.add_table(rows=len(entries), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for idx, (title, page_num, indent) in enumerate(entries):
        row = table.rows[idx]
        
        # Col 1: Title
        cell1 = row.cells[0]
        cell1.width = Inches(5.6)
        set_cell_border(cell1, val="none")
        set_cell_margins(cell1, top=20, bottom=20, start=0, end=0)
        p1 = cell1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.line_spacing = 1.15
        
        prefix = "  " * indent
        bold_flag = True if indent == 0 else False
        r1 = p1.add_run(prefix + title)
        set_font(r1, size=12 if indent == 0 else 11.5, bold=bold_flag)

        # Col 2: Page
        cell2 = row.cells[1]
        cell2.width = Inches(0.8)
        set_cell_border(cell2, val="none")
        set_cell_margins(cell2, top=20, bottom=20, start=0, end=0)
        p2 = cell2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2.paragraph_format.space_after = Pt(1)
        p2.paragraph_format.space_before = Pt(1)
        p2.paragraph_format.line_spacing = 1.15
        r2 = p2.add_run(str(page_num))
        set_font(r2, size=12 if indent == 0 else 11.5, bold=bold_flag)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(6)

print("Base setup ready.")
