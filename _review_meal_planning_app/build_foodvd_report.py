from __future__ import annotations

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
OUT = BASE / "Bao_cao_do_an_FoodVD.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color: str = "D9D9D9") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    values = {"top": top, "start": start, "bottom": bottom, "end": end}
    for key, value in values.items():
        tag = f"w:{key}"
        node = margins.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name: str = "Times New Roman", size: int | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_paragraph(doc: Document, text: str = "", style: str | None = None, bold_lead: str | None = None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.55
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if style is None else p.alignment
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead) :])
        set_font(rest)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_font(run, bold=True)
        run.font.color.rgb = RGBColor(0, 0, 0)
    if level == 1:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(8)
    else:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
    return p


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.3
        r = p.add_run(item)
        set_font(r)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.3
        r = p.add_run(item)
        set_font(r)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = header_cells[i]
        set_cell_shading(cell, "FFFFFF")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        set_font(run, size=10, bold=True)
        if widths:
            cell.width = Inches(widths[i])
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cell = cells[i]
            set_cell_shading(cell, "FFFFFF")
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if len(value) > 18 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(value)
            set_font(run, size=10)
            if widths:
                cell.width = Inches(widths[i])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    return table


def add_code(doc: Document, code: str) -> None:
    for line in code.strip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(line)
        set_font(run, name="Consolas", size=8)
    doc.add_paragraph()


def add_figure_placeholder(doc: Document, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"[Chèn hình: {caption}]")
    set_font(run, bold=True)
    p.paragraph_format.space_after = Pt(2)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = cap.add_run(caption)
    set_font(cap_run, size=10)
    cap_run.italic = True


PAGE_BREAK_COUNT = 0


def page_break(doc: Document) -> None:
    global PAGE_BREAK_COUNT
    PAGE_BREAK_COUNT += 1
    if PAGE_BREAK_COUNT <= 2:
        doc.add_page_break()
        return
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(6)
    spacer.paragraph_format.space_after = Pt(6)


def add_page_number(paragraph) -> None:
    prefix = paragraph.add_run("Trang ")
    set_font(prefix, size=9)

    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_char_begin)

    instr_run = paragraph.add_run()
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    instr_run._r.append(instr_text)

    separate_run = paragraph.add_run()
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    separate_run._r.append(fld_char_separate)

    number_run = paragraph.add_run("1")
    set_font(number_run, size=9)

    end_run = paragraph.add_run()
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    end_run._r.append(fld_char_end)


def enable_field_update_on_open(docx_path: Path) -> None:
    with ZipFile(docx_path, "r") as source:
        files = {name: source.read(name) for name in source.namelist()}

    settings_name = "word/settings.xml"
    settings = files.get(settings_name, b"").decode("utf-8")
    if settings and "w:updateFields" not in settings:
        settings = settings.replace(
            "</w:settings>",
            '<w:updateFields w:val="true"/></w:settings>',
        )
        files[settings_name] = settings.encode("utf-8")

    with ZipFile(docx_path, "w", ZIP_DEFLATED) as target:
        for name, data in files.items():
            target.writestr(name, data)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    styles["Normal"].font.size = Pt(12)
    for style_name, size in [("Title", 18), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
    return doc


def title_page(doc: Document) -> None:
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TRƯỜNG ĐẠI HỌC")
    set_font(r, size=13, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("KHOA CÔNG NGHỆ THÔNG TIN")
    set_font(r, size=13, bold=True)
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BÁO CÁO ĐỒ ÁN TỐT NGHIỆP")
    set_font(r, size=18, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("XÂY DỰNG HỆ THỐNG THƯƠNG MẠI ĐIỆN TỬ FOODVD NỀN TẢNG ĐẶT MÓN")
    set_font(r, size=15, bold=True)
    for _ in range(2):
        doc.add_paragraph()
    rows = [
        ["Sinh viên thực hiện", "Vũ Dũng"],
        ["Giảng viên hướng dẫn", "Bùi Thị Thanh"],
        ["Chuyên ngành", "Công nghệ thông tin / Kỹ thuật phần mềm"],
        ["Công nghệ chính", "React, TypeScript, Tailwind CSS, FastAPI, MongoDB"],
    ]
    add_table(doc, ["Thông tin", "Nội dung"], rows, widths=[2.2, 4.2])
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TP. Hồ Chí Minh, 2026")
    set_font(r, size=12, bold=True)
    page_break(doc)


def add_static_toc(doc: Document) -> None:
    add_heading(doc, "MỤC LỤC", 1)
    entries = [
        "Lời mở đầu",
        "Chương 1 Tổng quan đề tài",
        "Chương 2 Cơ sở lý thuyết",
        "Chương 3 Khảo sát bài toán thực tế",
        "Chương 4 Phân tích yêu cầu hệ thống",
        "Chương 5 Thiết kế hệ thống",
        "Chương 6 Công nghệ thư viện và môi trường triển khai",
        "Chương 7 Thuật toán và xử lý nghiệp vụ",
        "Chương 8 Thiết kế dữ liệu và bảo mật",
        "Chương 9 Thiết kế giao diện và triển khai prototype",
        "Chương 10 Kiểm thử đánh giá và hướng phát triển",
        "Phụ lục PlantUML",
    ]
    for i, entry in enumerate(entries, 1):
        p = doc.add_paragraph()
        r = p.add_run(f"{i}. {entry}")
        set_font(r)
    page_break(doc)


def intro(doc: Document) -> None:
    add_heading(doc, "LỜI MỞ ĐẦU", 1)
    paragraphs = [
        "Sự phát triển của thương mại điện tử đã thay đổi cách người tiêu dùng tiếp cận sản phẩm và dịch vụ hằng ngày. Trong lĩnh vực thực phẩm và đồ uống, nhu cầu đặt món trực tuyến tăng nhanh vì người dùng mong muốn tiết kiệm thời gian, chủ động lựa chọn món ăn, biết trước chi phí và theo dõi tiến trình giao hàng. Với các cửa hàng vừa và nhỏ, một nền tảng đặt món riêng cũng giúp giảm phụ thuộc vào các sàn lớn, kiểm soát dữ liệu khách hàng và chủ động vận hành kinh doanh.",
        "Đề tài FoodVD được xây dựng theo hướng một hệ thống thương mại điện tử chuyên cho đặt món ăn. Hệ thống hướng tới hai nhóm người dùng chính: khách hàng đặt món và quản trị viên hoặc chủ quán vận hành thực đơn, đơn hàng, khuyến mãi và báo cáo doanh thu. Bên cạnh yêu cầu nghiệp vụ, đề tài chú trọng các vấn đề kỹ thuật thực tế như thiết kế RESTful API, lưu trữ dữ liệu phi quan hệ bằng MongoDB, mã hóa dữ liệu nhạy cảm bằng AES-256, hỗ trợ thanh toán COD hoặc VietQR và gửi email hóa đơn sau khi tạo đơn.",
        "Báo cáo này trình bày toàn bộ quá trình phân tích và xây dựng hệ thống, từ cơ sở lý thuyết, khảo sát bài toán thực tế, đặc tả yêu cầu, thiết kế kiến trúc, mô hình dữ liệu, thuật toán nghiệp vụ, lựa chọn công nghệ, đến kiểm thử và hướng phát triển. Phần cuối báo cáo cung cấp mã PlantUML cho các sơ đồ use case, class, sequence, activity, component và deployment để người thực hiện có thể vẽ sơ đồ, xuất ảnh rồi chèn vào báo cáo.",
    ]
    for text in paragraphs:
        add_paragraph(doc, text)
    add_heading(doc, "Phạm vi bản báo cáo", 2)
    add_bullets(
        doc,
        [
            "Tập trung mô tả hệ thống FoodVD như một nền tảng đặt món trực tuyến cho cửa hàng hoặc chuỗi cửa hàng F&B.",
            "Mô tả frontend prototype đã triển khai bằng React, TypeScript, Vite và Tailwind CSS.",
            "Đề xuất backend theo FastAPI, MongoDB, AES-256, Gmail SMTP và VietQR để hoàn thiện hệ thống sản phẩm.",
            "Cung cấp các đoạn PlantUML để phục vụ vẽ sơ đồ trong quá trình hoàn thiện báo cáo.",
        ],
    )
    page_break(doc)


def figure_insertion_guide(doc: Document) -> None:
    add_heading(doc, "DANH MỤC HÌNH CẦN CHÈN", 1)
    add_paragraph(
        doc,
        "Các hình dưới đây là những hình nên có trong báo cáo để phần phân tích và thiết kế rõ ràng hơn. Bạn có thể dùng mã PlantUML ở phụ lục để xuất ảnh, sau đó chèn vào đúng vị trí được gợi ý. Khi chèn ảnh, nên căn giữa, đặt kích thước vừa chiều ngang trang và giữ caption ngay bên dưới hình.",
    )
    add_table(
        doc,
        ["STT", "Tên hình nên chèn", "Vị trí gợi ý", "Nguồn vẽ"],
        [
            ["1", "Use Case tổng quát hệ thống FoodVD", "Chương 4, sau mục tác nhân và yêu cầu chức năng.", "usecase_tong_quat.puml"],
            ["2", "Use Case khách hàng", "Chương 4, sau mô tả chức năng phía khách hàng.", "usecase_khach_hang.puml"],
            ["3", "Use Case quản trị viên", "Chương 4, sau mô tả chức năng admin.", "usecase_admin.puml"],
            ["4", "Kiến trúc tổng thể FoodVD", "Chương 5, sau mục kiến trúc tổng thể.", "component_foodvd.puml hoặc vẽ kiến trúc riêng"],
            ["5", "Class diagram hệ thống", "Chương 5 hoặc Chương 8, sau phần thiết kế dữ liệu.", "class_diagram_foodvd.puml"],
            ["6", "Sequence đặt hàng", "Chương 7, sau thuật toán tạo đơn hàng.", "sequence_dat_hang.puml"],
            ["7", "Activity đặt hàng", "Chương 7, sau mô tả quy trình nghiệp vụ.", "activity_dat_hang.puml"],
            ["8", "Deployment diagram", "Chương 6 hoặc Chương 5, sau môi trường triển khai.", "deployment_foodvd.puml"],
            ["9", "Ảnh giao diện thực đơn và giỏ hàng", "Chương 9, thay placeholder Hình 9.1.", "Chụp màn hình website"],
            ["10", "Ảnh giao diện quản trị đơn hàng", "Chương 9, thay placeholder Hình 9.2.", "Chụp màn hình website"],
        ],
        widths=[0.6, 2.3, 2.7, 1.4],
    )
    add_heading(doc, "Cách xuất hình từ PlantUML", 2)
    add_numbered(
        doc,
        [
            "Tạo thư mục diagrams trong project để chứa các file .puml.",
            "Copy từng đoạn mã PlantUML ở phụ lục vào file tương ứng, ví dụ usecase_tong_quat.puml.",
            "Chạy lệnh java -jar plantuml.jar ten_file.puml để xuất PNG.",
            "Mở Word, đặt con trỏ tại vị trí placeholder hình, chọn Insert, Pictures và chọn ảnh đã xuất.",
            "Căn giữa ảnh, đặt caption theo đúng số chương, ví dụ Hình 4.1 Use Case tổng quát hệ thống FoodVD.",
        ],
    )
    add_heading(doc, "Danh sách ảnh chụp giao diện nên bổ sung", 2)
    add_paragraph(
        doc,
        "Ngoài các sơ đồ PlantUML, báo cáo nên có một vài ảnh chụp màn hình website để chứng minh sản phẩm đã có giao diện và luồng thao tác. Các ảnh này không cần quá nhiều, nhưng nên chọn đúng những màn hình thể hiện chức năng chính của đề tài.",
    )
    add_table(
        doc,
        ["Ảnh", "Nội dung cần chụp", "Vị trí chèn"],
        [
            ["Ảnh 1", "Trang thực đơn có bộ lọc, danh sách món và giá bán.", "Chương 9, sau mục nguyên tắc thiết kế giao diện."],
            ["Ảnh 2", "Giỏ hàng có món, topping, voucher, tổng thanh toán và chọn VietQR.", "Chương 9, thay Hình 9.1 hoặc đặt ngay sau Hình 9.1."],
            ["Ảnh 3", "Mã VietQR demo hoặc khu vực thanh toán.", "Chương 7 hoặc Chương 9 khi mô tả thanh toán."],
            ["Ảnh 4", "Cổng quản trị đơn hàng có trạng thái Chờ duyệt, Đang chuẩn bị, Hoàn tất.", "Chương 9, thay Hình 9.2."],
            ["Ảnh 5", "Màn hình quản lý thực đơn, thêm món và bật tắt trạng thái còn hàng.", "Chương 9 hoặc phụ lục demo."],
            ["Ảnh 6", "Biểu đồ doanh thu và các chỉ số tổng đơn, đơn hủy, món bán chạy.", "Chương 10 hoặc phụ lục demo."],
        ],
        widths=[0.9, 3.5, 2.1],
    )
    page_break(doc)


def chapter_one(doc: Document) -> None:
    add_heading(doc, "CHƯƠNG 1 TỔNG QUAN ĐỀ TÀI", 1)
    add_heading(doc, "1.1 Lý do chọn đề tài", 2)
    for text in [
        "Trong bối cảnh chuyển đổi số, người tiêu dùng ngày càng quen với việc sử dụng website và ứng dụng để mua hàng, thanh toán và theo dõi dịch vụ. Ngành F&B là một trong những lĩnh vực chịu ảnh hưởng rõ rệt vì nhu cầu ăn uống diễn ra hằng ngày, tần suất đặt hàng cao và yêu cầu tốc độ phục vụ nhanh. Người dùng không chỉ cần xem món, mà còn cần biết tình trạng còn hàng, giá bán, topping, chương trình giảm giá, phí giao hàng và hình thức thanh toán phù hợp.",
        "Nhiều cửa hàng vừa và nhỏ vẫn xử lý đơn hàng qua điện thoại, tin nhắn hoặc sổ ghi chép. Cách làm này dễ dẫn đến sai sót khi tính tiền, bỏ sót topping, nhầm địa chỉ hoặc không kiểm soát được trạng thái đơn. Khi phụ thuộc hoàn toàn vào các nền tảng lớn, cửa hàng có thể chịu chiết khấu cao, khó xây dựng dữ liệu khách hàng riêng và ít chủ động trong chiến dịch khuyến mãi. Vì vậy, một hệ thống FoodVD riêng có ý nghĩa thực tiễn trong việc tối ưu quy trình đặt món và quản trị vận hành.",
        "Bên cạnh nghiệp vụ bán hàng, bài toán bảo vệ thông tin cá nhân cũng rất quan trọng. Số điện thoại, địa chỉ nhận hàng và lịch sử đơn hàng là dữ liệu nhạy cảm. Nếu hệ thống lưu trữ không an toàn, rủi ro rò rỉ dữ liệu có thể ảnh hưởng trực tiếp đến khách hàng và uy tín cửa hàng. Đề tài đưa vào cơ chế mã hóa AES-256 trước khi lưu dữ liệu nhạy cảm vào MongoDB nhằm thể hiện hướng tiếp cận bảo mật phù hợp với ứng dụng web hiện đại.",
    ]:
        add_paragraph(doc, text)
    add_heading(doc, "1.2 Mục tiêu đề tài", 2)
    add_bullets(
        doc,
        [
            "Xây dựng giao diện đặt món trực quan, hỗ trợ tìm kiếm, lọc món theo danh mục và giá bán.",
            "Cung cấp giỏ hàng có số lượng, topping, voucher, phí giao hàng và tổng thanh toán.",
            "Mô tả quy trình đặt hàng bằng COD hoặc VietQR, tạo mã đơn hàng và gửi email hóa đơn.",
            "Xây dựng cổng quản trị cho thực đơn, trạng thái món, xử lý đơn hàng và báo cáo doanh thu.",
            "Đề xuất kiến trúc backend FastAPI, cơ sở dữ liệu MongoDB và cơ chế mã hóa AES-256 cho dữ liệu nhạy cảm.",
        ],
    )
    add_heading(doc, "1.3 Đối tượng và phạm vi nghiên cứu", 2)
    add_table(
        doc,
        ["Nhóm", "Nội dung nghiên cứu", "Kết quả mong muốn"],
        [
            ["Khách hàng", "Xem thực đơn, lọc món, thêm topping, đặt hàng, thanh toán COD hoặc VietQR.", "Trải nghiệm đặt món nhanh, rõ tổng tiền, nhận thông báo đơn hàng."],
            ["Quản trị viên", "Quản lý món, danh mục, trạng thái còn hàng, voucher, đơn hàng và báo cáo.", "Vận hành cửa hàng chủ động, giảm thao tác thủ công."],
            ["Kỹ thuật", "React, FastAPI, MongoDB, AES-256, SMTP, VietQR và mô hình RESTful API.", "Hình thành kiến trúc có thể mở rộng thành sản phẩm thật."],
        ],
        widths=[1.4, 3.1, 2.0],
    )
    add_heading(doc, "1.4 Phương pháp thực hiện", 2)
    add_numbered(
        doc,
        [
            "Khảo sát quy trình đặt món phổ biến và xác định các tác nhân chính trong hệ thống.",
            "Phân tích yêu cầu chức năng, yêu cầu phi chức năng và ràng buộc bảo mật.",
            "Thiết kế kiến trúc tổng thể theo mô hình frontend, backend API và database.",
            "Xây dựng prototype frontend để minh họa luồng đặt món và quản trị.",
            "Đề xuất thuật toán xử lý nghiệp vụ, mô hình dữ liệu và kế hoạch kiểm thử.",
        ],
    )
    page_break(doc)


def chapter_two(doc: Document) -> None:
    add_heading(doc, "CHƯƠNG 2 CƠ SỞ LÝ THUYẾT", 1)
    sections = [
        (
            "2.1 Thương mại điện tử trong lĩnh vực F&B",
            [
                "Thương mại điện tử là mô hình giao dịch trong đó hoạt động tìm kiếm, đặt mua, thanh toán và chăm sóc khách hàng được thực hiện thông qua hệ thống số. Trong lĩnh vực F&B, thương mại điện tử có đặc thù khác với bán hàng hóa thông thường vì món ăn có thời gian chuẩn bị ngắn, thời hạn sử dụng ngắn và trạng thái còn hàng thay đổi liên tục. Do đó hệ thống cần phản hồi nhanh, giao diện rõ ràng và quy trình xử lý đơn phải phù hợp với vận hành bếp.",
                "Một nền tảng đặt món hiệu quả cần kết hợp ba lớp nghiệp vụ: lớp hiển thị thực đơn cho khách hàng, lớp xử lý đặt hàng và thanh toán, lớp quản trị vận hành cho cửa hàng. Nếu thiếu một trong ba lớp này, hệ thống dễ trở thành trang giới thiệu tĩnh hoặc công cụ ghi đơn đơn giản, chưa đủ đáp ứng nhu cầu thương mại điện tử hoàn chỉnh.",
            ],
        ),
        (
            "2.2 RESTful API và FastAPI",
            [
                "RESTful API là phong cách thiết kế dịch vụ web dựa trên tài nguyên. Mỗi nhóm dữ liệu như món ăn, danh mục, đơn hàng, voucher hoặc người dùng được xem như tài nguyên và được thao tác bằng các phương thức HTTP như GET, POST, PUT, PATCH và DELETE. Cách tiếp cận này giúp frontend và backend tách rời, dễ kiểm thử và dễ mở rộng.",
                "FastAPI là framework Python hiện đại cho phép xây dựng API nhanh, có kiểm tra kiểu dữ liệu bằng Pydantic và tự động sinh tài liệu OpenAPI. Với FoodVD, FastAPI phù hợp cho các endpoint như tìm kiếm món, tạo đơn hàng, cập nhật trạng thái, tạo voucher, thống kê doanh thu và gửi email hóa đơn.",
            ],
        ),
        (
            "2.3 MongoDB và dữ liệu phi quan hệ",
            [
                "MongoDB lưu dữ liệu theo dạng document JSON/BSON, phù hợp với các đối tượng có cấu trúc linh hoạt như món ăn có nhiều topping, đơn hàng có nhiều dòng sản phẩm, voucher có điều kiện áp dụng khác nhau. Thay vì phải tách nhiều bảng quan hệ nhỏ, hệ thống có thể lưu đơn hàng cùng danh sách items để truy xuất nhanh lịch sử đơn.",
                "Tuy nhiên, việc dùng MongoDB đòi hỏi thiết kế schema có kiểm soát. Các trường quan trọng như order_id, user_id, status, created_at, category và price cần được đánh chỉ mục để tăng tốc truy vấn. Với báo cáo doanh thu, MongoDB Aggregation Pipeline có thể nhóm đơn theo ngày hoặc tháng, tính tổng doanh thu và xác định món bán chạy.",
            ],
        ),
        (
            "2.4 Mã hóa AES-256",
            [
                "AES là thuật toán mã hóa đối xứng được sử dụng rộng rãi để bảo vệ dữ liệu nhạy cảm. AES-256 sử dụng khóa 256 bit, có độ an toàn cao nếu khóa được tạo ngẫu nhiên, lưu trữ riêng trong biến môi trường và không đưa trực tiếp vào mã nguồn. Trong FoodVD, số điện thoại và địa chỉ nhận hàng nên được mã hóa trước khi lưu vào MongoDB.",
                "Một thiết kế an toàn không chỉ mã hóa nội dung, mà còn phải sử dụng IV hoặc nonce ngẫu nhiên cho mỗi lần mã hóa, lưu kèm tag xác thực nếu dùng chế độ GCM, và giới hạn quyền truy cập khóa. Khi cần hiển thị cho admin, hệ thống có thể giải mã phía backend sau khi kiểm tra quyền, hoặc chỉ hiển thị dữ liệu đã che một phần để giảm rủi ro lộ thông tin.",
            ],
        ),
        (
            "2.5 SMTP và gửi email hóa đơn",
            [
                "SMTP là giao thức dùng để gửi email giữa ứng dụng và máy chủ thư. Với Gmail SMTP, hệ thống có thể gửi hóa đơn sau khi đơn hàng được tạo thành công. Nội dung email cần có mã đơn hàng, danh sách món, tổng tiền, hình thức thanh toán và thông tin hỗ trợ. Trong môi trường thật, tài khoản gửi mail cần dùng app password hoặc cơ chế OAuth phù hợp.",
                "Gửi email nên được xử lý bất đồng bộ hoặc đưa vào hàng đợi nếu hệ thống có lưu lượng lớn. Việc này giúp API tạo đơn không bị chậm do phụ thuộc phản hồi của máy chủ email. Nếu gửi thất bại, đơn hàng vẫn có thể được tạo, sau đó hệ thống ghi log và thử gửi lại.",
            ],
        ),
        (
            "2.6 VietQR trong thanh toán",
            [
                "VietQR là chuẩn mã QR chuyển khoản ngân hàng phổ biến tại Việt Nam. Khi khách chọn thanh toán VietQR, hệ thống có thể sinh mã QR chứa thông tin ngân hàng, số tài khoản, số tiền và nội dung chuyển khoản. Với prototype, phần VietQR có thể mô phỏng bằng mã QR demo; khi triển khai thật cần tích hợp thư viện hoặc API tạo QR theo chuẩn.",
                "Để tự động xác nhận thanh toán, hệ thống cần kết nối webhook hoặc dịch vụ kiểm tra biến động tài khoản. Nếu chưa có đối soát tự động, admin có thể xác nhận thủ công khi nhận được giao dịch. Báo cáo cần phân biệt rõ giữa mô phỏng giao diện và tích hợp thanh toán thật.",
            ],
        ),
    ]
    for heading, paragraphs in sections:
        add_heading(doc, heading, 2)
        for text in paragraphs:
            add_paragraph(doc, text)
    add_table(
        doc,
        ["Khái niệm", "Vai trò trong FoodVD", "Lưu ý triển khai"],
        [
            ["RESTful API", "Giao tiếp giữa frontend và backend.", "Chuẩn hóa mã lỗi, phân quyền và validation."],
            ["MongoDB", "Lưu món ăn, đơn hàng, voucher, người dùng.", "Thiết kế index và schema rõ ràng."],
            ["AES-256", "Bảo vệ SĐT và địa chỉ nhận hàng.", "Không hard-code khóa, dùng IV ngẫu nhiên."],
            ["SMTP", "Gửi hóa đơn tự động.", "Xử lý bất đồng bộ và ghi log lỗi."],
            ["VietQR", "Thanh toán chuyển khoản nhanh.", "Cần đối soát giao dịch khi triển khai thật."],
        ],
        widths=[1.3, 2.8, 2.4],
    )
    page_break(doc)


def chapter_three(doc: Document) -> None:
    add_heading(doc, "CHƯƠNG 3 KHẢO SÁT BÀI TOÁN THỰC TẾ", 1)
    add_heading(doc, "3.1 Hiện trạng đặt món tại cửa hàng vừa và nhỏ", 2)
    for text in [
        "Nhiều cửa hàng ăn uống quy mô vừa và nhỏ tiếp nhận đơn qua điện thoại, tin nhắn mạng xã hội hoặc nhân viên ghi trực tiếp. Quy trình này đơn giản khi số lượng đơn ít, nhưng trở nên khó kiểm soát khi đơn tăng vào giờ cao điểm. Nhân viên có thể ghi thiếu topping, nhầm món, không cập nhật kịp món hết hàng hoặc tính sai khuyến mãi.",
        "Một vấn đề khác là dữ liệu khách hàng bị phân tán. Số điện thoại, địa chỉ, lịch sử mua hàng và phản hồi khách có thể nằm ở nhiều kênh khác nhau. Cửa hàng khó phân tích món bán chạy, khung giờ cao điểm và hiệu quả voucher. Nếu có hệ thống riêng, dữ liệu được tập trung và có thể dùng cho quản trị doanh thu, chăm sóc khách hàng và cải thiện thực đơn.",
    ]:
        add_paragraph(doc, text)
    add_heading(doc, "3.2 Các vấn đề cần giải quyết", 2)
    add_table(
        doc,
        ["Vấn đề", "Tác động", "Hướng giải quyết trong FoodVD"],
        [
            ["Tìm món mất thời gian", "Khách khó chọn nhanh, giảm tỷ lệ đặt hàng.", "Tìm kiếm theo tên, lọc danh mục và khoảng giá."],
            ["Sai sót khi ghi đơn", "Nhầm số lượng, topping, địa chỉ.", "Giỏ hàng rõ từng dòng, form thông tin và xác nhận tổng tiền."],
            ["Món hết nhưng vẫn nhận đơn", "Phải gọi lại khách, trải nghiệm kém.", "Admin bật tắt trạng thái món theo thời gian thực."],
            ["Không kiểm soát voucher", "Áp dụng sai điều kiện, thất thoát doanh thu.", "Kiểm tra min order, loại giảm giá và thời hạn."],
            ["Rò rỉ thông tin cá nhân", "Ảnh hưởng khách hàng và uy tín cửa hàng.", "Mã hóa SĐT, địa chỉ bằng AES-256 trước khi lưu."],
            ["Thiếu báo cáo", "Khó ra quyết định kinh doanh.", "Aggregation theo ngày, tháng, món bán chạy, đơn hủy."],
        ],
        widths=[1.6, 2.2, 2.7],
    )
    add_heading(doc, "3.3 Quy trình nghiệp vụ đề xuất", 2)
    add_numbered(
        doc,
        [
            "Khách truy cập website FoodVD và xem danh sách món đang mở bán.",
            "Khách tìm kiếm hoặc lọc món theo danh mục, khoảng giá và nhu cầu cá nhân.",
            "Khách thêm món vào giỏ, chọn topping, số lượng và áp dụng voucher nếu có.",
            "Khách nhập thông tin nhận hàng và chọn phương thức thanh toán COD hoặc VietQR.",
            "Hệ thống kiểm tra dữ liệu, tính tổng tiền, mã hóa thông tin nhạy cảm và tạo đơn hàng.",
            "Nếu chọn VietQR, hệ thống sinh mã QR hoặc nội dung chuyển khoản để khách thanh toán.",
            "Hệ thống gửi email hóa đơn, admin tiếp nhận và cập nhật trạng thái đơn.",
            "Sau khi hoàn tất, dữ liệu đơn hàng được dùng cho báo cáo doanh thu và món bán chạy.",
        ],
    )
    add_figure_placeholder(doc, "Hình 3.1 Quy trình đặt món và xử lý đơn hàng của FoodVD")
    page_break(doc)


def chapter_four(doc: Document) -> None:
    add_heading(doc, "CHƯƠNG 4 PHÂN TÍCH YÊU CẦU HỆ THỐNG", 1)
    add_heading(doc, "4.1 Tác nhân hệ thống", 2)
    add_table(
        doc,
        ["Tác nhân", "Mô tả", "Quyền chính"],
        [
            ["Khách vãng lai", "Người chưa đăng nhập hoặc chưa có tài khoản.", "Xem thực đơn, tìm kiếm, lọc món, thêm giỏ tạm."],
            ["Khách thành viên", "Người có tài khoản và lịch sử mua hàng.", "Đặt hàng, theo dõi đơn, hủy đơn chờ duyệt, nhận hóa đơn."],
            ["Quản trị viên", "Chủ quán hoặc nhân viên quản lý hệ thống.", "CRUD món, voucher, đơn hàng, báo cáo, trạng thái mở bán."],
            ["Cổng thanh toán", "Dịch vụ tạo mã VietQR hoặc xác nhận giao dịch.", "Sinh QR, đối soát thanh toán."],
            ["Dịch vụ email", "Gmail SMTP hoặc dịch vụ gửi thư.", "Gửi hóa đơn và thông báo đơn hàng."],
        ],
        widths=[1.5, 2.3, 2.7],
    )
    add_heading(doc, "4.2 Yêu cầu chức năng", 2)
    requirements = [
        ["F01", "Tìm kiếm và lọc món", "Khách có thể nhập từ khóa, chọn danh mục và khoảng giá để lọc thực đơn."],
        ["F02", "Quản lý giỏ hàng", "Thêm món, tăng giảm số lượng, chọn topping, tính tạm tính."],
        ["F03", "Áp dụng voucher", "Kiểm tra điều kiện voucher và tính giảm giá trước khi đặt hàng."],
        ["F04", "Tạo đơn hàng", "Lưu thông tin đơn, khách hàng, items, thanh toán, ghi chú và trạng thái ban đầu."],
        ["F05", "Thanh toán COD/VietQR", "Cho phép khách chọn COD hoặc hiển thị mã QR chuyển khoản."],
        ["F06", "Gửi email hóa đơn", "Gửi email cho khách sau khi tạo đơn thành công."],
        ["F07", "Quản lý thực đơn", "Admin thêm, sửa, xóa món, cập nhật ảnh, giá, topping và trạng thái còn hàng."],
        ["F08", "Xử lý đơn hàng", "Admin tiếp nhận, cập nhật trạng thái, hủy đơn khi cần."],
        ["F09", "Báo cáo doanh thu", "Thống kê doanh thu theo thời gian, tổng đơn, đơn hủy, món bán chạy."],
    ]
    add_table(doc, ["Mã", "Chức năng", "Mô tả"], requirements, widths=[0.8, 1.9, 3.8])
    add_heading(doc, "4.3 Yêu cầu phi chức năng", 2)
    add_bullets(
        doc,
        [
            "Hiệu năng: thao tác tìm kiếm và lọc món phản hồi nhanh, API danh sách món có phân trang khi dữ liệu lớn.",
            "Bảo mật: mật khẩu băm bằng thuật toán an toàn, SĐT và địa chỉ mã hóa AES-256, API admin yêu cầu JWT và phân quyền.",
            "Khả dụng: giao diện tương thích desktop, tablet và điện thoại; các thao tác chính không phụ thuộc vào màn hình lớn.",
            "Dễ bảo trì: frontend tách component, backend tách router, service, repository và schema validation.",
            "Khả năng mở rộng: có thể bổ sung nhiều chi nhánh, nhiều cửa hàng, tích hợp thanh toán và giao hàng bên thứ ba.",
        ],
    )
    add_figure_placeholder(doc, "Hình 4.1 Use Case tổng quát hệ thống FoodVD")
    page_break(doc)


def chapter_five(doc: Document) -> None:
    add_heading(doc, "CHƯƠNG 5 THIẾT KẾ HỆ THỐNG", 1)
    add_heading(doc, "5.1 Kiến trúc tổng thể", 2)
    for text in [
        "FoodVD được thiết kế theo kiến trúc ba lớp gồm frontend, backend API và cơ sở dữ liệu. Frontend chịu trách nhiệm hiển thị giao diện, thu thập thao tác người dùng, kiểm tra dữ liệu cơ bản và gọi API. Backend FastAPI xử lý nghiệp vụ, xác thực, phân quyền, mã hóa dữ liệu, tạo đơn, gửi email và truy vấn MongoDB. MongoDB lưu dữ liệu người dùng, thực đơn, voucher, đơn hàng và log thanh toán.",
        "Kiến trúc tách lớp giúp mỗi thành phần có trách nhiệm rõ ràng. Frontend có thể được triển khai bằng Vite/React và phân phối qua web server hoặc CDN. Backend có thể triển khai độc lập, mở rộng theo số lượng request. Database có thể cấu hình index, backup và phân quyền riêng để bảo vệ dữ liệu.",
    ]:
        add_paragraph(doc, text)
    add_figure_placeholder(doc, "Hình 5.1 Kiến trúc tổng thể FoodVD")
    add_heading(doc, "5.2 Thiết kế module", 2)
    add_table(
        doc,
        ["Module", "Chức năng", "Thành phần chính"],
        [
            ["Catalog", "Hiển thị thực đơn, tìm kiếm, lọc món.", "Dish API, category, price filter, search index."],
            ["Cart", "Quản lý món đã chọn, topping, số lượng.", "LocalStorage/session, cart service, price calculator."],
            ["Order", "Tạo đơn, cập nhật trạng thái, hủy đơn.", "Order API, state machine, email event."],
            ["Payment", "COD, VietQR, nội dung chuyển khoản.", "VietQR generator, payment status, reconciliation."],
            ["Admin", "Quản trị món, voucher, đơn hàng, báo cáo.", "Admin routes, role-based access control."],
            ["Security", "Mã hóa và phân quyền.", "AES service, JWT, password hashing, environment secrets."],
        ],
        widths=[1.2, 2.6, 2.7],
    )
    add_heading(doc, "5.3 Thiết kế API đề xuất", 2)
    add_table(
        doc,
        ["Phương thức", "Endpoint", "Mục đích"],
        [
            ["GET", "/api/dishes", "Lấy danh sách món theo keyword, category, price range."],
            ["POST", "/api/orders", "Tạo đơn hàng mới và kích hoạt gửi email hóa đơn."],
            ["GET", "/api/orders/{id}", "Theo dõi chi tiết và trạng thái đơn hàng."],
            ["PATCH", "/api/admin/orders/{id}/status", "Admin cập nhật trạng thái đơn."],
            ["POST", "/api/admin/dishes", "Thêm món ăn mới."],
            ["PATCH", "/api/admin/dishes/{id}", "Sửa món, giá, ảnh, topping, trạng thái còn hàng."],
            ["GET", "/api/admin/reports/revenue", "Thống kê doanh thu theo khoảng thời gian."],
            ["POST", "/api/payments/vietqr", "Sinh mã VietQR hoặc nội dung chuyển khoản."],
        ],
        widths=[1.0, 2.3, 3.2],
    )
    add_heading(doc, "5.4 Luồng trạng thái đơn hàng", 2)
    add_paragraph(
        doc,
        "Đơn hàng đi qua các trạng thái chính gồm Chờ duyệt, Đang chuẩn bị, Đang giao, Hoàn tất và Đã hủy. Chỉ những đơn ở trạng thái Chờ duyệt mới được khách hủy. Admin có thể chuyển trạng thái theo luồng hợp lệ; hệ thống cần ghi lại thời gian cập nhật để phục vụ truy vết và báo cáo.",
    )
    add_table(
        doc,
        ["Trạng thái hiện tại", "Hành động", "Trạng thái tiếp theo"],
        [
            ["Chờ duyệt", "Admin tiếp nhận", "Đang chuẩn bị"],
            ["Chờ duyệt", "Khách/Admin hủy", "Đã hủy"],
            ["Đang chuẩn bị", "Hoàn tất chế biến", "Đang giao"],
            ["Đang giao", "Giao thành công", "Hoàn tất"],
            ["Hoàn tất", "Không cập nhật", "Hoàn tất"],
        ],
        widths=[1.8, 2.2, 2.0],
    )
    page_break(doc)


def chapter_six(doc: Document) -> None:
    add_heading(doc, "CHƯƠNG 6 CÔNG NGHỆ THƯ VIỆN VÀ MÔI TRƯỜNG TRIỂN KHAI", 1)
    add_heading(doc, "6.1 Frontend", 2)
    add_table(
        doc,
        ["Công nghệ", "Vai trò", "Lý do lựa chọn"],
        [
            ["React", "Xây dựng giao diện theo component.", "Phù hợp UI tương tác như catalog, giỏ hàng, admin dashboard."],
            ["TypeScript", "Kiểm tra kiểu dữ liệu.", "Giảm lỗi khi thao tác order, dish, voucher và trạng thái."],
            ["Vite", "Dev server và bundler.", "Khởi động nhanh, build gọn cho React."],
            ["Tailwind CSS", "Thiết kế giao diện responsive.", "Tạo layout, màu sắc và spacing nhanh, nhất quán."],
            ["oxfmt", "Định dạng mã nguồn.", "Hỗ trợ thống nhất style trong source."],
        ],
        widths=[1.4, 2.2, 2.9],
    )
    add_paragraph(
        doc,
        "Prototype hiện tại đã triển khai frontend bằng React, TypeScript, Vite và Tailwind CSS. Giao diện gồm thực đơn, bộ lọc, giỏ hàng, checkout COD/VietQR, cổng quản trị đơn hàng, quản lý thực đơn và biểu đồ doanh thu demo. Đây là phần minh họa trực quan cho đề tài và có thể kết nối backend thật ở giai đoạn tiếp theo.",
    )
    add_heading(doc, "6.2 Backend đề xuất", 2)
    add_table(
        doc,
        ["Thành phần", "Thư viện/Framework", "Ghi chú triển khai"],
        [
            ["API", "FastAPI", "Tách router theo auth, dishes, orders, admin, reports."],
            ["Validation", "Pydantic", "Định nghĩa schema request/response, kiểm tra dữ liệu đầu vào."],
            ["Database driver", "Motor hoặc PyMongo", "Motor phù hợp async API; PyMongo phù hợp xử lý đồng bộ."],
            ["Authentication", "python-jose, passlib", "JWT access token, refresh token, hash mật khẩu."],
            ["Encryption", "cryptography", "AES-256-GCM cho SĐT và địa chỉ."],
            ["Email", "smtplib hoặc aiosmtplib", "Gửi hóa đơn qua Gmail SMTP."],
            ["QR", "qrcode hoặc API VietQR", "Sinh QR demo hoặc tích hợp chuẩn VietQR thật."],
        ],
        widths=[1.5, 2.0, 3.0],
    )
    add_heading(doc, "6.3 Môi trường phát triển", 2)
    add_bullets(
        doc,
        [
            "Visual Studio Code dùng để viết mã nguồn frontend và backend.",
            "Git/GitHub dùng để quản lý phiên bản, tạo nhánh chức năng và lưu lịch sử thay đổi.",
            "MongoDB Compass hoặc MongoDB Atlas dùng để kiểm tra dữ liệu trong quá trình phát triển.",
            "Postman hoặc Swagger UI dùng để kiểm thử RESTful API.",
            "PlantUML dùng để vẽ use case, class, sequence, activity, component và deployment diagram.",
        ],
    )
    add_heading(doc, "6.4 Cấu trúc thư mục đề xuất", 2)
    add_code(
        doc,
        """
foodvd/
  frontend/
    src/
      components/
      pages/
      services/
      types/
  backend/
    app/
      routers/
      services/
      repositories/
      schemas/
      core/
    tests/
  docs/
    diagrams/
    report/
""",
    )
    page_break(doc)


def chapter_seven(doc: Document) -> None:
    add_heading(doc, "CHƯƠNG 7 THUẬT TOÁN VÀ XỬ LÝ NGHIỆP VỤ", 1)
    add_heading(doc, "7.1 Thuật toán tìm kiếm và lọc món", 2)
    add_paragraph(
        doc,
        "Tìm kiếm món cần kết hợp từ khóa, danh mục và khoảng giá. Với frontend prototype, dữ liệu mẫu được lọc trực tiếp trên mảng dish. Với backend thật, API nhận keyword, category, min_price và max_price rồi tạo truy vấn MongoDB. Keyword có thể dùng biểu thức chính quy không phân biệt hoa thường; category và price dùng điều kiện chính xác hoặc so sánh.",
    )
    add_code(
        doc,
        """
Input: keyword, category, min_price, max_price
query = {}
if keyword is not empty:
    query[\"name\"] = {\"$regex\": keyword, \"$options\": \"i\"}
if category != \"all\":
    query[\"category\"] = category
query[\"price\"] = {\"$gte\": min_price, \"$lte\": max_price}
result = db.dishes.find(query).sort({\"sold\": -1})
Output: list of dishes
""",
    )
    add_heading(doc, "7.2 Thuật toán tính giỏ hàng", 2)
    add_paragraph(
        doc,
        "Giỏ hàng tính tổng theo từng dòng món. Mỗi dòng gồm món, số lượng và danh sách topping. Đơn giá thực tế bằng giá khuyến mãi nếu có, ngược lại dùng giá niêm yết, cộng thêm chi phí topping. Tổng đơn bằng tổng các dòng, trừ voucher hợp lệ và cộng phí giao hàng nếu chưa đạt ngưỡng miễn phí.",
    )
    add_table(
        doc,
        ["Bước", "Mô tả", "Kết quả"],
        [
            ["1", "Lấy đơn giá món.", "promo_price hoặc price."],
            ["2", "Tính phí topping.", "Số topping nhân đơn giá topping."],
            ["3", "Tính tiền dòng.", "(đơn giá + topping) nhân số lượng."],
            ["4", "Tính voucher.", "Giảm theo phần trăm hoặc số tiền cố định."],
            ["5", "Tính tổng cuối.", "subtotal - discount + shipping_fee."],
        ],
        widths=[0.8, 3.0, 2.7],
    )
    add_heading(doc, "7.3 Thuật toán kiểm tra voucher", 2)
    add_paragraph(
        doc,
        "Voucher cần kiểm tra nhiều điều kiện trước khi áp dụng. Các điều kiện thường gồm mã voucher tồn tại, còn hiệu lực, chưa vượt số lượt dùng, đơn hàng đạt giá trị tối thiểu, người dùng thuộc nhóm được áp dụng và món trong giỏ không thuộc danh sách loại trừ. Việc kiểm tra phải thực hiện ở backend để tránh người dùng chỉnh sửa logic phía frontend.",
    )
    add_code(
        doc,
        """
if voucher is None:
    discount = 0
elif current_time not between voucher.start_at and voucher.end_at:
    discount = 0
elif subtotal < voucher.min_order_value:
    discount = 0
elif voucher.used_count >= voucher.usage_limit:
    discount = 0
else:
    discount = calculate_discount(voucher, subtotal)
""",
    )
    add_heading(doc, "7.4 Thuật toán tạo đơn hàng", 2)
    add_numbered(
        doc,
        [
            "Nhận request gồm danh sách món, topping, số lượng, thông tin khách, voucher và phương thức thanh toán.",
            "Kiểm tra dữ liệu bắt buộc, định dạng email, số điện thoại, địa chỉ và sự tồn tại của từng món.",
            "Kiểm tra trạng thái còn hàng, giá tại thời điểm đặt và điều kiện voucher.",
            "Tính tổng tiền cuối cùng và sinh mã đơn hàng duy nhất.",
            "Mã hóa SĐT và địa chỉ bằng AES-256-GCM trước khi lưu.",
            "Tạo bản ghi order trong MongoDB với trạng thái Chờ duyệt.",
            "Nếu thanh toán VietQR, sinh thông tin QR hoặc nội dung chuyển khoản.",
            "Kích hoạt tác vụ gửi email hóa đơn và trả kết quả cho frontend.",
        ],
    )
    add_heading(doc, "7.5 Thuật toán báo cáo doanh thu", 2)
    add_paragraph(
        doc,
        "Báo cáo doanh thu chỉ tính các đơn ở trạng thái Hoàn tất. Dữ liệu được lọc theo khoảng thời gian, sau đó nhóm theo ngày hoặc tháng. MongoDB Aggregation Pipeline phù hợp cho bài toán này vì có thể lọc, bung mảng items, nhóm tổng tiền và sắp xếp kết quả ngay trong database.",
    )
    add_code(
        doc,
        """
db.orders.aggregate([
  { $match: { status: \"Hoàn tất\", created_at: { $gte: from, $lte: to } } },
  { $group: {
      _id: { $dateToString: { format: \"%Y-%m-%d\", date: \"$created_at\" } },
      revenue: { $sum: \"$total\" },
      order_count: { $sum: 1 }
  }},
  { $sort: { _id: 1 } }
])
""",
    )
    add_heading(doc, "7.6 Mã hóa AES-256-GCM", 2)
    add_paragraph(
        doc,
        "AES-256-GCM vừa mã hóa vừa cung cấp tag xác thực để phát hiện dữ liệu bị sửa đổi. Mỗi bản ghi nên có IV ngẫu nhiên, ciphertext và tag. Khóa bí mật không lưu trong database mà đặt trong biến môi trường hoặc secret manager. Khi cần tìm kiếm theo SĐT, có thể lưu thêm trường hash một chiều của SĐT đã chuẩn hóa để tra cứu mà không cần giải mã hàng loạt.",
    )
    add_code(
        doc,
        """
def encrypt_sensitive_text(plain_text, key):
    iv = random_bytes(12)
    cipher = AESGCM(key)
    cipher_text = cipher.encrypt(iv, plain_text.encode(), None)
    return base64(iv + cipher_text)

def decrypt_sensitive_text(payload, key):
    raw = base64_decode(payload)
    iv = raw[:12]
    cipher_text = raw[12:]
    return AESGCM(key).decrypt(iv, cipher_text, None).decode()
""",
    )
    page_break(doc)


def chapter_eight(doc: Document) -> None:
    add_heading(doc, "CHƯƠNG 8 THIẾT KẾ DỮ LIỆU VÀ BẢO MẬT", 1)
    add_heading(doc, "8.1 Các collection chính", 2)
    add_table(
        doc,
        ["Collection", "Trường tiêu biểu", "Mục đích"],
        [
            ["users", "name, email, password_hash, role, phone_encrypted, address_encrypted", "Lưu tài khoản khách hàng và admin."],
            ["categories", "name, description, status", "Quản lý danh mục món."],
            ["dishes", "name, category_id, price, promo_price, image, toppings, available", "Lưu thực đơn và trạng thái mở bán."],
            ["vouchers", "code, type, value, min_order, start_at, end_at, usage_limit", "Quản lý khuyến mãi."],
            ["orders", "order_id, user_id, items, total, payment_method, status, encrypted_contact", "Lưu đơn hàng."],
            ["payment_logs", "order_id, provider, amount, status, transaction_ref", "Theo dõi thanh toán."],
            ["audit_logs", "actor_id, action, target, created_at", "Ghi lịch sử thao tác admin."],
        ],
        widths=[1.4, 3.3, 1.8],
    )
    add_heading(doc, "8.2 Schema đơn hàng đề xuất", 2)
    add_code(
        doc,
        """
{
  \"order_id\": \"FVD-20260914-0001\",
  \"customer_id\": \"ObjectId or null\",
  \"items\": [
    { \"dish_id\": \"ObjectId\", \"name\": \"Phở bò\", \"price\": 69000,
      \"quantity\": 2, \"toppings\": [\"Quẩy\"] }
  ],
  \"subtotal\": 145000,
  \"discount\": 20000,
  \"shipping_fee\": 15000,
  \"total\": 140000,
  \"payment_method\": \"VietQR\",
  \"payment_status\": \"pending\",
  \"status\": \"Chờ duyệt\",
  \"phone_encrypted\": \"...\",
  \"address_encrypted\": \"...\",
  \"created_at\": \"2026-09-14T10:30:00Z\"
}
""",
    )
    add_heading(doc, "8.3 Chỉ mục dữ liệu", 2)
    add_table(
        doc,
        ["Collection", "Index", "Lý do"],
        [
            ["dishes", "name text hoặc regex-support, category_id, price", "Tăng tốc tìm kiếm và lọc món."],
            ["orders", "order_id unique", "Tra cứu đơn hàng nhanh và tránh trùng mã."],
            ["orders", "status, created_at", "Lọc đơn theo trạng thái và báo cáo thời gian."],
            ["vouchers", "code unique", "Kiểm tra voucher nhanh."],
            ["users", "email unique", "Đăng nhập và chống trùng tài khoản."],
            ["orders", "phone_hash", "Tra cứu đơn theo SĐT mà không lưu SĐT rõ."],
        ],
        widths=[1.4, 2.5, 2.6],
    )
    add_heading(doc, "8.4 Bảo mật hệ thống", 2)
    add_bullets(
        doc,
        [
            "Xác thực người dùng bằng JWT, phân biệt role khách hàng và admin.",
            "Băm mật khẩu bằng bcrypt hoặc argon2, không lưu mật khẩu dạng rõ.",
            "Mã hóa SĐT và địa chỉ bằng AES-256-GCM trước khi ghi database.",
            "Dùng HTTPS trong môi trường triển khai thật để bảo vệ dữ liệu trên đường truyền.",
            "Giới hạn CORS theo domain frontend chính thức.",
            "Ghi audit log cho thao tác admin như sửa giá, tắt món, hủy đơn và cập nhật trạng thái.",
            "Không đưa khóa AES, mật khẩu email, chuỗi kết nối MongoDB vào mã nguồn hoặc repository.",
        ],
    )
    add_heading(doc, "8.5 Rủi ro và biện pháp giảm thiểu", 2)
    add_table(
        doc,
        ["Rủi ro", "Mức độ", "Biện pháp"],
        [
            ["Lộ khóa mã hóa", "Cao", "Lưu secret trong biến môi trường, xoay khóa định kỳ, giới hạn quyền truy cập."],
            ["Spam tạo đơn ảo", "Trung bình", "Rate limit, captcha nhẹ khi bất thường, xác nhận email hoặc số điện thoại."],
            ["Áp voucher sai", "Trung bình", "Kiểm tra voucher ở backend, ghi lịch sử sử dụng."],
            ["Sai lệch thanh toán QR", "Cao", "Đối soát giao dịch, trạng thái payment_status, không hoàn tất đơn nếu chưa xác nhận."],
            ["Admin thao tác nhầm", "Trung bình", "Xác nhận hành động quan trọng, audit log, phân quyền rõ."],
        ],
        widths=[2.0, 1.2, 3.3],
    )
    page_break(doc)


def chapter_nine(doc: Document) -> None:
    add_heading(doc, "CHƯƠNG 9 THIẾT KẾ GIAO DIỆN VÀ TRIỂN KHAI PROTOTYPE", 1)
    add_heading(doc, "9.1 Nguyên tắc thiết kế giao diện", 2)
    for text in [
        "Giao diện FoodVD được định hướng như một bề mặt làm việc thực tế, không phải chỉ là landing page giới thiệu. Màn hình đầu tiên cần cho người dùng thấy ngay thực đơn, khả năng tìm kiếm, bộ lọc và giỏ hàng. Nhờ đó khách có thể bắt đầu đặt món mà không cần đi qua nhiều lớp thông tin quảng cáo.",
        "Màu đỏ được dùng làm màu nhận diện chính vì phù hợp với lĩnh vực ẩm thực và gợi cảm giác năng lượng. Nền sáng giúp thực đơn dễ đọc, còn khu vực quản trị dùng nền tối để phân tách rõ giữa trải nghiệm khách hàng và trải nghiệm vận hành. Các thẻ món cần có ảnh, tên món, giá, mô tả, trạng thái còn hàng và nút thêm vào giỏ.",
    ]:
        add_paragraph(doc, text)
    add_heading(doc, "9.2 Các màn hình chính", 2)
    add_table(
        doc,
        ["Màn hình", "Nội dung", "Mục tiêu"],
        [
            ["Trang thực đơn", "Danh sách món, bộ lọc, tìm kiếm, giá tối đa.", "Giúp khách tìm món nhanh."],
            ["Giỏ hàng", "Món đã chọn, topping, số lượng, voucher, tổng tiền.", "Minh bạch chi phí trước khi đặt."],
            ["Checkout", "Thông tin khách, phương thức COD/VietQR, QR demo.", "Tạo đơn hàng rõ ràng."],
            ["Admin đơn hàng", "Danh sách đơn, trạng thái, nút cập nhật/hủy.", "Điều phối vận hành."],
            ["Admin thực đơn", "Thêm món, bật/tắt trạng thái còn hàng.", "Quản lý món theo thời gian thực."],
            ["Báo cáo", "Doanh thu, tổng đơn, đơn hủy, món bán chạy, biểu đồ.", "Hỗ trợ quyết định kinh doanh."],
        ],
        widths=[1.5, 3.0, 2.0],
    )
    add_figure_placeholder(doc, "Hình 9.1 Giao diện thực đơn và giỏ hàng FoodVD")
    add_figure_placeholder(doc, "Hình 9.2 Giao diện cổng quản trị đơn hàng")
    add_heading(doc, "9.3 Prototype frontend đã triển khai", 2)
    add_paragraph(
        doc,
        "Prototype hiện tại được triển khai trong project React/Vite. Các kiểu dữ liệu chính gồm Category, PaymentMethod, OrderStatus, Dish, CartItem và Order. Ứng dụng sử dụng state React để lưu thực đơn mẫu, giỏ hàng, voucher, phương thức thanh toán, form khách hàng, đơn hàng và tab admin. Các hàm nghiệp vụ chính gồm addToCart, changeQuantity, placeOrder, advanceOrder và addDishFromAdmin.",
    )
    add_table(
        doc,
        ["Hàm", "Vai trò", "Ý nghĩa nghiệp vụ"],
        [
            ["addToCart", "Thêm món vào giỏ, cộng số lượng nếu trùng.", "Mô phỏng thao tác chọn món/topping."],
            ["changeQuantity", "Tăng giảm số lượng và xóa dòng khi về 0.", "Quản lý giỏ hàng tạm."],
            ["placeOrder", "Tạo đơn, che SĐT/địa chỉ, thêm vào danh sách đơn.", "Mô phỏng tạo đơn và gửi hóa đơn."],
            ["advanceOrder", "Chuyển trạng thái đơn theo luồng.", "Mô phỏng xử lý đơn hàng của admin."],
            ["addDishFromAdmin", "Thêm món mới từ form admin.", "Mô phỏng CRUD thực đơn."],
        ],
        widths=[1.7, 2.5, 2.3],
    )
    add_heading(doc, "9.4 Hướng kết nối backend thật", 2)
    add_bullets(
        doc,
        [
            "Tách dữ liệu món và đơn hàng khỏi state mẫu, thay bằng API service gọi FastAPI.",
            "Dùng LocalStorage cho giỏ hàng khách vãng lai hoặc lưu giỏ hàng server-side cho thành viên.",
            "Khi bấm tạo đơn, frontend gửi payload đến POST /api/orders thay vì tự tạo đơn trong state.",
            "Admin dashboard lấy dữ liệu từ endpoint admin và cần JWT role admin.",
            "Báo cáo doanh thu gọi API aggregation thay vì tính trên mảng demo.",
        ],
    )
    page_break(doc)


def chapter_ten(doc: Document) -> None:
    add_heading(doc, "CHƯƠNG 10 KIỂM THỬ ĐÁNH GIÁ VÀ HƯỚNG PHÁT TRIỂN", 1)
    add_heading(doc, "10.1 Kiểm thử chức năng", 2)
    add_table(
        doc,
        ["Mã test", "Kịch bản", "Kết quả mong đợi"],
        [
            ["TC01", "Tìm món theo từ khóa phở.", "Danh sách chỉ hiển thị món phù hợp."],
            ["TC02", "Lọc danh mục Đồ uống.", "Chỉ hiển thị món thuộc danh mục đồ uống."],
            ["TC03", "Thêm món có topping vào giỏ.", "Giỏ hàng có đúng món, topping và số lượng."],
            ["TC04", "Áp voucher FOODVD20 khi đủ điều kiện.", "Tổng tiền giảm 20% theo quy định."],
            ["TC05", "Tạo đơn VietQR.", "Đơn mới ở trạng thái Chờ duyệt và hiển thị thông tin QR."],
            ["TC06", "Admin cập nhật trạng thái.", "Đơn chuyển theo luồng Chờ duyệt đến Hoàn tất."],
            ["TC07", "Admin tắt món còn hàng.", "Món chuyển trạng thái Tạm hết và không thêm được vào giỏ."],
            ["TC08", "Xem báo cáo doanh thu.", "Hiển thị tổng doanh thu, số đơn, đơn hủy và biểu đồ."],
        ],
        widths=[0.9, 3.1, 2.5],
    )
    add_heading(doc, "10.2 Kiểm thử bảo mật", 2)
    add_bullets(
        doc,
        [
            "Kiểm tra API admin không cho truy cập nếu thiếu token hoặc token không có role admin.",
            "Kiểm tra dữ liệu SĐT và địa chỉ trong MongoDB không lưu dạng rõ.",
            "Kiểm tra khóa AES, mật khẩu Gmail SMTP và chuỗi kết nối MongoDB không xuất hiện trong mã nguồn.",
            "Kiểm tra giới hạn request tạo đơn để giảm spam.",
            "Kiểm tra validation dữ liệu đầu vào để tránh giá âm, số lượng âm, email sai định dạng hoặc voucher không hợp lệ.",
        ],
    )
    add_heading(doc, "10.3 Đánh giá kết quả", 2)
    for text in [
        "Prototype đã thể hiện được luồng nghiệp vụ chính của nền tảng đặt món: tìm kiếm, lọc món, giỏ hàng, voucher, COD/VietQR, tạo đơn, quản trị đơn hàng, quản lý thực đơn và báo cáo. Giao diện bám sát nhu cầu thực tế của khách hàng và admin, không chỉ dừng ở trang giới thiệu.",
        "Tuy nhiên, để trở thành hệ thống hoàn chỉnh, cần xây dựng backend FastAPI, database MongoDB, xác thực người dùng, mã hóa AES-256 thật, gửi email hóa đơn thật và tích hợp VietQR theo chuẩn sản xuất. Đây là các phần có thể phát triển tiếp dựa trên kiến trúc và đặc tả đã trình bày trong báo cáo.",
    ]:
        add_paragraph(doc, text)
    add_heading(doc, "10.4 Hướng phát triển", 2)
    add_numbered(
        doc,
        [
            "Hoàn thiện API backend với FastAPI, JWT, Pydantic và MongoDB.",
            "Tích hợp đăng ký, đăng nhập, quản lý tài khoản và lịch sử đơn hàng.",
            "Tích hợp VietQR thật và đối soát thanh toán tự động.",
            "Tách dashboard admin thành khu vực riêng có phân quyền.",
            "Bổ sung thông báo thời gian thực bằng WebSocket cho đơn hàng mới.",
            "Tối ưu báo cáo bằng biểu đồ chi tiết theo ngày, tháng, danh mục và món bán chạy.",
            "Bổ sung kiểm thử tự động frontend, backend và kiểm thử bảo mật cơ bản.",
        ],
    )
    add_heading(doc, "KẾT LUẬN", 1)
    add_paragraph(
        doc,
        "Đề tài FoodVD có tính thực tiễn cao vì giải quyết nhu cầu đặt món trực tuyến và quản trị vận hành cho cửa hàng F&B. Báo cáo đã trình bày từ lý thuyết nền tảng, bài toán thực tế, yêu cầu hệ thống, thiết kế kiến trúc, công nghệ, thuật toán, bảo mật đến kiểm thử. Prototype frontend hiện tại là cơ sở trực quan để phát triển tiếp backend và hoàn thiện thành hệ thống thương mại điện tử đặt món đầy đủ.",
    )
    page_break(doc)


def appendix_implementation_plan(doc: Document) -> None:
    add_heading(doc, "PHỤ LỤC KẾ HOẠCH TRIỂN KHAI BACKEND", 1)
    add_paragraph(
        doc,
        "Phụ lục này dùng để định hướng giai đoạn phát triển tiếp theo sau prototype frontend. Mục tiêu là chuyển các chức năng demo trong giao diện thành hệ thống có backend thật, database thật và cơ chế bảo mật có thể kiểm thử được. Kế hoạch được chia thành các giai đoạn nhỏ để dễ quản lý tiến độ đồ án.",
    )
    add_table(
        doc,
        ["Giai đoạn", "Nội dung thực hiện", "Kết quả bàn giao"],
        [
            ["1", "Khởi tạo FastAPI, cấu hình CORS, cấu trúc router, service, repository và schema.", "Backend chạy được, Swagger UI hiển thị endpoint mẫu."],
            ["2", "Kết nối MongoDB, tạo collection dishes, categories, vouchers, orders và users.", "CRUD món ăn và danh mục hoạt động."],
            ["3", "Xây dựng đăng ký, đăng nhập, JWT, role khách hàng và admin.", "API admin được bảo vệ bằng phân quyền."],
            ["4", "Xây dựng API đặt hàng, tính giỏ hàng, voucher, phí giao hàng và trạng thái đơn.", "Frontend có thể tạo đơn qua API."],
            ["5", "Tích hợp AES-256-GCM cho SĐT và địa chỉ nhận hàng.", "Dữ liệu nhạy cảm trong MongoDB không lưu dạng rõ."],
            ["6", "Tích hợp Gmail SMTP gửi hóa đơn sau khi tạo đơn.", "Khách nhận được email hóa đơn."],
            ["7", "Tích hợp VietQR demo hoặc API VietQR thật.", "Đơn VietQR có QR và nội dung chuyển khoản."],
            ["8", "Xây dựng báo cáo doanh thu bằng Aggregation Pipeline.", "Admin xem doanh thu, đơn hủy và món bán chạy."],
        ],
        widths=[0.9, 3.6, 2.0],
    )
    add_heading(doc, "Checklist nghiệm thu backend", 2)
    add_bullets(
        doc,
        [
            "Tất cả endpoint có schema request và response rõ ràng, có ví dụ trong Swagger.",
            "Các endpoint admin từ chối request thiếu token hoặc token không có quyền admin.",
            "Đơn hàng lưu đúng items, số lượng, topping, tổng tiền, voucher và trạng thái.",
            "Dữ liệu phone_encrypted và address_encrypted trong database không đọc được trực tiếp.",
            "Email hóa đơn có mã đơn, danh sách món, tổng tiền và phương thức thanh toán.",
            "Báo cáo doanh thu chỉ tính đơn Hoàn tất, không tính đơn Đã hủy hoặc Chờ duyệt.",
        ],
    )
    page_break(doc)


def appendix_api_contract(doc: Document) -> None:
    add_heading(doc, "PHỤ LỤC ĐẶC TẢ API MẪU", 1)
    add_paragraph(
        doc,
        "Các API dưới đây là hợp đồng dữ liệu mẫu để kết nối frontend hiện tại với backend FastAPI. Khi triển khai thật, tên trường nên giữ ổn định để giao diện không phải thay đổi nhiều. Những trường nhạy cảm như phone và address chỉ nhận ở request tạo đơn; backend sẽ mã hóa trước khi lưu.",
    )
    add_table(
        doc,
        ["API", "Request chính", "Response chính"],
        [
            ["GET /api/dishes", "keyword, category, min_price, max_price, page", "items, total, page, page_size"],
            ["POST /api/orders", "items, customer, voucher_code, payment_method, note", "order_id, total, status, vietqr_payload"],
            ["GET /api/orders/{id}", "order_id", "order detail, status timeline"],
            ["PATCH /api/admin/orders/{id}/status", "status", "updated order"],
            ["POST /api/admin/dishes", "name, category, price, image, toppings", "created dish"],
            ["GET /api/admin/reports/revenue", "from_date, to_date, group_by", "revenue series, total, best sellers"],
        ],
        widths=[1.9, 2.5, 2.1],
    )
    add_heading(doc, "Payload tạo đơn mẫu", 2)
    add_code(
        doc,
        """
{
  \"items\": [
    { \"dish_id\": \"65f...\", \"quantity\": 2, \"toppings\": [\"Quẩy\"] },
    { \"dish_id\": \"66a...\", \"quantity\": 1, \"toppings\": [\"Ít đá\"] }
  ],
  \"customer\": {
    \"name\": \"Vũ Dũng\",
    \"phone\": \"0901234567\",
    \"email\": \"dung.foodvd@example.com\",
    \"address\": \"12 Nguyễn Huệ, Quận 1, TP.HCM\"
  },
  \"voucher_code\": \"FOODVD20\",
  \"payment_method\": \"VietQR\",
  \"note\": \"Ít hành, giao giờ trưa\"
}
""",
    )
    add_heading(doc, "Response tạo đơn mẫu", 2)
    add_code(
        doc,
        """
{
  \"order_id\": \"FVD-20260914-0001\",
  \"status\": \"Chờ duyệt\",
  \"subtotal\": 167000,
  \"discount\": 33400,
  \"shipping_fee\": 0,
  \"total\": 133600,
  \"payment_method\": \"VietQR\",
  \"vietqr_payload\": \"bank=...&amount=133600&content=FVD-20260914-0001\",
  \"message\": \"Đặt hàng thành công. Hóa đơn đã được gửi qua email.\"
}
""",
    )
    page_break(doc)


def appendix_traceability(doc: Document) -> None:
    add_heading(doc, "PHỤ LỤC MA TRẬN TRUY VẾT YÊU CẦU", 1)
    add_paragraph(
        doc,
        "Ma trận truy vết giúp chứng minh rằng mỗi yêu cầu chức năng đều có use case, API hoặc màn hình tương ứng và có kịch bản kiểm thử. Đây là phần nên có trong báo cáo đồ án vì giúp hội đồng thấy hệ thống được thiết kế có kiểm soát, không chỉ mô tả rời rạc từng chức năng.",
    )
    add_table(
        doc,
        ["Mã yêu cầu", "Use case liên quan", "API hoặc màn hình", "Test case"],
        [
            ["F01", "Tìm kiếm và lọc món", "GET /api/dishes, màn hình Thực đơn", "TC01, TC02"],
            ["F02", "Quản lý giỏ hàng", "Cart UI, LocalStorage hoặc cart service", "TC03"],
            ["F03", "Áp dụng voucher", "Voucher service, checkout panel", "TC04"],
            ["F04", "Tạo đơn hàng", "POST /api/orders, checkout form", "TC05"],
            ["F05", "Thanh toán COD/VietQR", "Payment module, QR panel", "TC05"],
            ["F06", "Gửi email hóa đơn", "Email service, Gmail SMTP", "TC05, kiểm thử email"],
            ["F07", "Quản lý thực đơn", "Admin menu screen, dish API", "TC07"],
            ["F08", "Xử lý đơn hàng", "Admin order screen, status API", "TC06"],
            ["F09", "Báo cáo doanh thu", "Report API, dashboard chart", "TC08"],
        ],
        widths=[1.0, 2.2, 2.5, 1.3],
    )
    add_heading(doc, "Quy tắc đánh giá hoàn thành yêu cầu", 2)
    add_bullets(
        doc,
        [
            "Một yêu cầu được xem là hoàn thành khi có giao diện thao tác, xử lý nghiệp vụ và dữ liệu trả về đúng.",
            "Yêu cầu liên quan đến bảo mật chỉ được xem là hoàn thành khi kiểm tra được dữ liệu nhạy cảm không lưu dạng rõ.",
            "Yêu cầu báo cáo chỉ được xem là hoàn thành khi dữ liệu lọc đúng trạng thái đơn Hoàn tất và đúng khoảng thời gian.",
            "Yêu cầu thanh toán VietQR trong prototype được chấp nhận ở mức mô phỏng; khi triển khai thật cần có đối soát giao dịch.",
        ],
    )
    page_break(doc)


def appendix_demo_script(doc: Document) -> None:
    add_heading(doc, "PHỤ LỤC KỊCH BẢN DEMO BẢO VỆ", 1)
    add_paragraph(
        doc,
        "Kịch bản demo nên ngắn gọn, đi theo một câu chuyện từ phía khách hàng sang phía quản trị viên. Người trình bày không nên mở quá nhiều màn hình cùng lúc, mà nên chuẩn bị sẵn dữ liệu mẫu và trình bày đúng các chức năng nổi bật của đề tài.",
    )
    add_table(
        doc,
        ["Bước", "Thao tác demo", "Ý nghĩa cần nói khi bảo vệ"],
        [
            ["1", "Mở trang FoodVD và giới thiệu thực đơn.", "Hệ thống tập trung vào nghiệp vụ đặt món, không chỉ là trang giới thiệu."],
            ["2", "Tìm kiếm món và lọc theo danh mục hoặc giá.", "Minh họa chức năng tìm kiếm, lọc dữ liệu và cải thiện trải nghiệm khách hàng."],
            ["3", "Thêm món vào giỏ, chọn topping và đổi số lượng.", "Chứng minh giỏ hàng xử lý được biến thể món và số lượng."],
            ["4", "Nhập thông tin khách, áp voucher và chọn VietQR.", "Trình bày cách hệ thống tính tổng tiền và hỗ trợ thanh toán."],
            ["5", "Tạo đơn hàng.", "Giải thích backend thật sẽ mã hóa SĐT, địa chỉ và gửi email hóa đơn."],
            ["6", "Chuyển sang cổng quản trị đơn hàng.", "Admin có thể tiếp nhận và cập nhật trạng thái đơn theo quy trình vận hành."],
            ["7", "Bật tắt trạng thái món hoặc thêm món mới.", "Chủ quán có thể kiểm soát món còn hàng theo thời gian thực."],
            ["8", "Mở phần báo cáo doanh thu.", "Dữ liệu đơn hoàn tất được dùng để thống kê doanh thu và món bán chạy."],
        ],
        widths=[0.7, 2.7, 3.1],
    )
    add_heading(doc, "Câu trả lời gợi ý khi hội đồng hỏi", 2)
    add_table(
        doc,
        ["Câu hỏi", "Ý trả lời ngắn gọn"],
        [
            ["Vì sao chọn MongoDB?", "Đơn hàng có items và topping linh hoạt, MongoDB lưu document phù hợp và aggregation hỗ trợ báo cáo."],
            ["Vì sao cần AES-256?", "SĐT và địa chỉ là dữ liệu nhạy cảm, cần mã hóa trước khi lưu database để giảm rủi ro rò rỉ."],
            ["VietQR đã thanh toán thật chưa?", "Prototype mô phỏng QR; bản triển khai thật cần API VietQR và cơ chế đối soát giao dịch."],
            ["Frontend hiện có kết nối backend chưa?", "Frontend hiện là prototype; báo cáo đã đặc tả API để kết nối FastAPI ở giai đoạn tiếp theo."],
            ["Hệ thống mở rộng như thế nào?", "Có thể tách module theo frontend, API, database, thêm WebSocket, phân quyền và báo cáo nâng cao."],
        ],
        widths=[2.3, 4.2],
    )
    page_break(doc)


def appendix_plantuml(doc: Document) -> None:
    add_heading(doc, "PHỤ LỤC PLANTUML", 1)
    add_paragraph(
        doc,
        "Bạn có thể copy từng đoạn mã dưới đây vào file .puml, sau đó dùng PlantUML để xuất ảnh PNG hoặc SVG rồi chèn vào các vị trí hình trong báo cáo. Nếu dùng PlantUML dạng jar, câu lệnh cơ bản là:",
    )
    add_code(
        doc,
        """
java -jar plantuml.jar usecase_tong_quat.puml
java -jar plantuml.jar usecase_khach_hang.puml
java -jar plantuml.jar usecase_admin.puml
java -jar plantuml.jar class_diagram_foodvd.puml
java -jar plantuml.jar sequence_dat_hang.puml
java -jar plantuml.jar activity_dat_hang.puml
java -jar plantuml.jar component_foodvd.puml
java -jar plantuml.jar deployment_foodvd.puml
""",
    )
    diagrams = [
        (
            "Use Case tổng quát",
            """
@startuml
left to right direction
actor \"Khách vãng lai\" as Guest
actor \"Khách thành viên\" as Member
actor \"Quản trị viên\" as Admin
actor \"Dịch vụ VietQR\" as VietQR
actor \"Gmail SMTP\" as SMTP

rectangle \"Hệ thống FoodVD\" {
  usecase \"Xem thực đơn\" as UC1
  usecase \"Tìm kiếm và lọc món\" as UC2
  usecase \"Quản lý giỏ hàng\" as UC3
  usecase \"Đặt hàng\" as UC4
  usecase \"Thanh toán COD\" as UC5
  usecase \"Thanh toán VietQR\" as UC6
  usecase \"Theo dõi đơn hàng\" as UC7
  usecase \"Hủy đơn chờ duyệt\" as UC8
  usecase \"Quản lý thực đơn\" as UC9
  usecase \"Quản lý voucher\" as UC10
  usecase \"Xử lý đơn hàng\" as UC11
  usecase \"Xem báo cáo doanh thu\" as UC12
  usecase \"Gửi email hóa đơn\" as UC13
}

Guest --> UC1
Guest --> UC2
Guest --> UC3
Guest --> UC4
Member --> UC7
Member --> UC8
UC4 --> UC5 : <<extend>>
UC4 --> UC6 : <<extend>>
UC4 --> UC13 : <<include>>
UC6 --> VietQR
UC13 --> SMTP
Admin --> UC9
Admin --> UC10
Admin --> UC11
Admin --> UC12
@enduml
""",
        ),
        (
            "Use Case khách hàng",
            """
@startuml
left to right direction
actor \"Khách hàng\" as Customer
rectangle \"FoodVD Customer Site\" {
  usecase \"Xem danh sách món\" as A
  usecase \"Xem chi tiết món\" as B
  usecase \"Lọc theo danh mục\" as C
  usecase \"Lọc theo giá\" as D
  usecase \"Chọn topping\" as E
  usecase \"Thêm vào giỏ\" as F
  usecase \"Áp dụng voucher\" as G
  usecase \"Nhập thông tin nhận hàng\" as H
  usecase \"Chọn COD hoặc VietQR\" as I
  usecase \"Xác nhận đặt hàng\" as J
}
Customer --> A
Customer --> B
Customer --> C
Customer --> D
Customer --> E
Customer --> F
Customer --> G
Customer --> H
Customer --> I
Customer --> J
J .> G : <<include>>
J .> H : <<include>>
J .> I : <<include>>
@enduml
""",
        ),
        (
            "Use Case quản trị viên",
            """
@startuml
left to right direction
actor \"Quản trị viên\" as Admin
rectangle \"FoodVD Admin\" {
  usecase \"Đăng nhập admin\" as A
  usecase \"Thêm món\" as B
  usecase \"Sửa món\" as C
  usecase \"Xóa món\" as D
  usecase \"Bật tắt trạng thái món\" as E
  usecase \"Tạo voucher\" as F
  usecase \"Tiếp nhận đơn mới\" as G
  usecase \"Cập nhật trạng thái đơn\" as H
  usecase \"Hủy đơn\" as I
  usecase \"Xem doanh thu\" as J
  usecase \"Xem món bán chạy\" as K
}
Admin --> A
Admin --> B
Admin --> C
Admin --> D
Admin --> E
Admin --> F
Admin --> G
Admin --> H
Admin --> I
Admin --> J
Admin --> K
@enduml
""",
        ),
        (
            "Class diagram",
            """
@startuml
class User {
  +id: ObjectId
  +name: string
  +email: string
  +passwordHash: string
  +role: string
}

class Dish {
  +id: ObjectId
  +name: string
  +categoryId: ObjectId
  +price: number
  +promoPrice: number
  +available: boolean
}

class Category {
  +id: ObjectId
  +name: string
}

class CartItem {
  +dishId: ObjectId
  +quantity: number
  +toppings: string[]
}

class Order {
  +id: ObjectId
  +orderCode: string
  +status: string
  +paymentMethod: string
  +total: number
}

class Voucher {
  +code: string
  +type: string
  +value: number
  +minOrder: number
}

class PaymentLog {
  +provider: string
  +amount: number
  +status: string
}

Category \"1\" -- \"*\" Dish
User \"1\" -- \"*\" Order
Order \"1\" -- \"*\" CartItem
Dish \"1\" -- \"*\" CartItem
Order \"0..1\" -- Voucher
Order \"1\" -- \"0..*\" PaymentLog
@enduml
""",
        ),
        (
            "Sequence đặt hàng",
            """
@startuml
actor Customer as C
participant \"React Frontend\" as FE
participant \"FastAPI Backend\" as API
database \"MongoDB\" as DB
participant \"AES Service\" as AES
participant \"VietQR Service\" as QR
participant \"Gmail SMTP\" as SMTP

C -> FE: Chọn món và bấm đặt hàng
FE -> API: POST /api/orders
API -> DB: Kiểm tra món và voucher
DB --> API: Dữ liệu hợp lệ
API -> AES: Mã hóa SĐT và địa chỉ
AES --> API: Ciphertext
API -> DB: Lưu đơn hàng
DB --> API: order_id
alt Thanh toán VietQR
  API -> QR: Tạo mã QR
  QR --> API: QR payload
end
API -> SMTP: Gửi email hóa đơn
SMTP --> API: Kết quả gửi
API --> FE: Thông tin đơn hàng
FE --> C: Hiển thị đặt hàng thành công
@enduml
""",
        ),
        (
            "Activity đặt hàng",
            """
@startuml
start
:Khách xem thực đơn;
:Tìm kiếm hoặc lọc món;
:Thêm món và topping vào giỏ;
if (Có voucher?) then (Có)
  :Kiểm tra điều kiện voucher;
else (Không)
endif
:Nhập thông tin nhận hàng;
:Chọn COD hoặc VietQR;
:Tính tổng tiền;
:Mã hóa SĐT và địa chỉ;
:Tạo đơn hàng trạng thái Chờ duyệt;
if (VietQR?) then (Có)
  :Sinh mã QR chuyển khoản;
else (COD)
endif
:Gửi email hóa đơn;
:Admin tiếp nhận đơn;
stop
@enduml
""",
        ),
        (
            "Component diagram",
            """
@startuml
package \"Frontend\" {
  [React App]
  [Cart UI]
  [Admin Dashboard]
}
package \"Backend\" {
  [FastAPI Routers]
  [Order Service]
  [Dish Service]
  [Voucher Service]
  [AES Service]
  [Email Service]
  [Report Service]
}
database \"MongoDB\" as DB
cloud \"Gmail SMTP\" as SMTP
cloud \"VietQR Provider\" as QR

[React App] --> [FastAPI Routers]
[Cart UI] --> [FastAPI Routers]
[Admin Dashboard] --> [FastAPI Routers]
[FastAPI Routers] --> [Order Service]
[FastAPI Routers] --> [Dish Service]
[FastAPI Routers] --> [Voucher Service]
[Order Service] --> [AES Service]
[Order Service] --> [Email Service]
[Order Service] --> QR
[Email Service] --> SMTP
[Report Service] --> DB
[Dish Service] --> DB
[Voucher Service] --> DB
[Order Service] --> DB
@enduml
""",
        ),
        (
            "Deployment diagram",
            """
@startuml
node \"Client Browser\" {
  artifact \"FoodVD React UI\"
}
node \"Web Server/CDN\" {
  artifact \"Static Build\"
}
node \"API Server\" {
  artifact \"FastAPI App\"
}
database \"MongoDB Server\" {
  artifact \"FoodVD Database\"
}
cloud \"Gmail SMTP\"
cloud \"VietQR API\"

\"Client Browser\" --> \"Web Server/CDN\" : HTTPS
\"Client Browser\" --> \"API Server\" : REST API HTTPS
\"API Server\" --> \"MongoDB Server\" : MongoDB driver
\"API Server\" --> \"Gmail SMTP\" : SMTP TLS
\"API Server\" --> \"VietQR API\" : HTTPS
@enduml
""",
        ),
    ]
    for title, code in diagrams:
        add_heading(doc, title, 2)
        add_code(doc, code)


def references(doc: Document) -> None:
    add_heading(doc, "TÀI LIỆU THAM KHẢO", 1)
    refs = [
        "FastAPI Documentation, phần xây dựng REST API và Pydantic schema.",
        "MongoDB Manual, phần document model, indexing và aggregation pipeline.",
        "React Documentation, phần component, state và event handling.",
        "Vite Documentation, phần development server và production build.",
        "Tailwind CSS Documentation, phần utility-first styling và responsive design.",
        "Python cryptography documentation, phần AESGCM.",
        "PlantUML Documentation, phần use case, sequence, class, activity, component và deployment diagram.",
    ]
    add_numbered(doc, refs)
    add_heading(doc, "Ghi chú hoàn thiện báo cáo", 2)
    add_paragraph(
        doc,
        "Khi chèn ảnh sơ đồ vào báo cáo, nên đặt ảnh ngay dưới đoạn mô tả tương ứng và giữ caption thống nhất theo dạng Hình x.y. Các hình use case nên đặt ở Chương 4, các hình kiến trúc và component nên đặt ở Chương 5, các hình sequence và activity nên đặt ở Chương 7. Nếu ảnh quá rộng, nên xuất SVG hoặc PNG độ phân giải cao rồi căn giữa theo chiều ngang trang.",
    )
    add_paragraph(
        doc,
        "Trước khi nộp bản cuối, người thực hiện nên cập nhật lại mục lục tự động trong Microsoft Word, kiểm tra số trang, rà chính tả tiếng Việt, thay các placeholder hình bằng ảnh thật và bổ sung tên trường, khoa, lớp, mã sinh viên nếu nhà trường yêu cầu theo mẫu riêng.",
    )
    add_heading(doc, "Checklist nộp đồ án", 2)
    add_table(
        doc,
        ["Hạng mục", "Việc cần kiểm tra", "Trạng thái"],
        [
            ["Báo cáo", "Đủ bìa, mục lục, chương nội dung, kết luận, tài liệu tham khảo và phụ lục.", "Chưa đánh dấu"],
            ["Sơ đồ", "Đã xuất ảnh PlantUML rõ nét và chèn đúng vị trí có caption.", "Chưa đánh dấu"],
            ["Source code", "Frontend build thành công, không lỗi TypeScript, cấu trúc thư mục rõ ràng.", "Chưa đánh dấu"],
            ["Demo", "Chuẩn bị kịch bản tìm món, đặt hàng, VietQR demo, admin cập nhật trạng thái.", "Chưa đánh dấu"],
            ["Backend", "Nếu có triển khai, API chạy được trong Swagger và kết nối MongoDB.", "Chưa đánh dấu"],
            ["Bảo mật", "Không để mật khẩu, khóa AES, chuỗi MongoDB hoặc app password trong source.", "Chưa đánh dấu"],
            ["Slide", "Slide ngắn gọn, có mục tiêu, kiến trúc, demo, kết quả và hướng phát triển.", "Chưa đánh dấu"],
        ],
        widths=[1.4, 3.8, 1.3],
    )


def main() -> None:
    doc = setup_document()
    title_page(doc)
    add_static_toc(doc)
    intro(doc)
    figure_insertion_guide(doc)
    chapter_one(doc)
    chapter_two(doc)
    chapter_three(doc)
    chapter_four(doc)
    chapter_five(doc)
    chapter_six(doc)
    chapter_seven(doc)
    chapter_eight(doc)
    chapter_nine(doc)
    chapter_ten(doc)
    appendix_implementation_plan(doc)
    appendix_api_contract(doc)
    appendix_traceability(doc)
    appendix_demo_script(doc)
    appendix_plantuml(doc)
    references(doc)

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_page_number(footer)

    doc.save(OUT)
    enable_field_update_on_open(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
