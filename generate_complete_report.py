# -*- coding: utf-8 -*-
"""
HỆ THỐNG XUẤT BÁO CÁO ĐỒ ÁN KỲ CHUẨN MỰC
Đề tài: XÂY DỰNG WEBSITE THƯƠNG MẠI ĐIỆN TỬ BÁN ĐỒ ĂN TRỰC TUYẾN FOODVD
        HỖ TRỢ TRẢI NGHIỆM 3D VÀ QUẢN LÝ VẬN HÀNH ĐA PHÂN HỆ
"""

from __future__ import annotations
import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

BASE = Path(r"c:\Users\Admin\Desktop\523100B\DOANTOTNGHIEP")
DIAGRAMS = BASE / "diagrams"
OUT_PATH = BASE / "BAO_CAO_DO_AN_KY_FOODVD_HOAN_THIEN.docx"

doc = Document()

# Configure Page Setup (A4, Left 3cm, Right 2cm, Top 2cm, Bottom 2cm)
for sec in doc.sections:
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)

# Helpers
def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_border(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)

def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    values = {"top": top, "start": start, "bottom": bottom, "end": end}
    for key, value in values.items():
        tag = f"w:{key}"
        node = margins.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_font(run, name="Times New Roman", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if color:
        run.font.color.rgb = color
    else:
        run.font.color.rgb = RGBColor(0, 0, 0)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic

def p(text="", bold_lead=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, italic=False):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(space_after)
    par.paragraph_format.line_spacing = 1.35
    par.alignment = align
    if bold_lead and text.startswith(bold_lead):
        r1 = par.add_run(bold_lead)
        set_font(r1, bold=True, size=13)
        r2 = par.add_run(text[len(bold_lead):])
        set_font(r2, size=13, italic=italic)
    else:
        r = par.add_run(text)
        set_font(r, size=13, italic=italic)
    return par

def h(text, level=1):
    par = doc.add_heading(text, level=level)
    par.paragraph_format.keep_with_next = True
    par.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    par.paragraph_format.space_after = Pt(6)
    size = 15 if level == 1 else (13.5 if level == 2 else 13)
    color = RGBColor(15, 23, 42)
    for r in par.runs:
        set_font(r, bold=True, size=size, color=color)
    return par

def bullets(items: list[str]):
    for it in items:
        par = doc.add_paragraph(style="List Bullet")
        par.paragraph_format.space_after = Pt(3)
        par.paragraph_format.line_spacing = 1.3
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = par.add_run(it)
        set_font(r, size=13)

def fig(filename, caption, width=Inches(5.8)):
    path = DIAGRAMS / filename
    if path.exists():
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(8)
        par.paragraph_format.space_after = Pt(4)
        run = par.add_run()
        run.add_picture(str(path), width=width)

        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(2)
        cap.paragraph_format.space_after = Pt(10)
        rc = cap.add_run(caption)
        set_font(rc, size=11, bold=True, italic=True, color=RGBColor(51, 65, 85))
    else:
        print(f"Figure not found: {filename}")

def tbl(headers, rows, widths=None):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header
    for i, title in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = title
        set_cell_shading(cell, "1E293B")
        set_cell_margins(cell, top=100, bottom=100, start=100, end=100)
        set_cell_border(cell, "0F172A")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        par = cell.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in par.runs:
            set_font(r, size=10.5, bold=True, color=RGBColor(255, 255, 255))

    # Data
    for r_idx, r_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cell = row.cells[c_idx]
            cell.text = str(val)
            set_cell_shading(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, start=90, end=90)
            set_cell_border(cell, "CBD5E1")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            par = cell.paragraphs[0]
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            if c_idx == 0:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in par.runs:
                set_font(r, size=10)

    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

print("Setup completed. Building document sections...")
